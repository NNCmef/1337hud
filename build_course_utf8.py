from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches, Pt as PPTPt
from pptx.dml.color import RGBColor

OUT = Path(r"C:\Users\Administrator\Downloads")
ASSETS = OUT / "1337_assets"

def add_p(doc, text="", center=False, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0 if center else 1.25)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    return p

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = value
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)

def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(14)

    for text in [
        "Министерство просвещения Республики Башкортостан",
        "ГАПОУ «Уфимский колледж статистики, информатики и вычислительной техники»",
        "", "КУРСОВОЙ ПРОЕКТ",
        "Разработка вычислительной локальной сети офиса «Ксанакс» с системой мониторинга технических узлов на базе отечественного программного комплекса «1337 Control»",
        "Пояснительная записка", "", "Студент группы 23СА-2: Калачев Роман Дмитриевич",
        "Руководитель: Поглазов Константин Юрьевич", "", "Уфа, 2026"
    ]:
        add_p(doc, text, center=True, bold=text == "КУРСОВОЙ ПРОЕКТ")
    doc.add_page_break()

    sections = [
        ("ЗАДАНИЕ", ["Спроектировать локальную вычислительную сеть офиса «Ксанакс» на 30 рабочих станций и 5 серверов. Выполнить VLAN-сегментацию, составить IPv4-план, определить серверные роли, построить модель в Cisco Packet Tracer и интегрировать программный комплекс мониторинга 1337 Control."]),
        ("АННОТАЦИЯ", ["В курсовом проекте разработана локальная вычислительная сеть офиса «Ксанакс», включающая 30 рабочих станций и пять серверов. Рассмотрены топология, VLAN, адресация, серверный сегмент, эмуляция Cisco Packet Tracer и отечественный программный комплекс 1337 Control.", "Ключевые слова: ЛВС, VLAN, Cisco Packet Tracer, мониторинг, FastAPI, SQLite, SwiftUI."]),
        ("ВВЕДЕНИЕ", ["Современный офис зависит от устойчивой работы вычислительной сети. Локальная сеть обеспечивает доступ сотрудников к документам, прикладным системам, печати, интернету и серверным ресурсам.", "Цель проекта — разработать ЛВС офиса «Ксанакс» на 30 рабочих станций и 5 серверов, смоделировать её в Cisco Packet Tracer и подключить 1337 Control.", "Задачи: анализ объекта, выбор топологии, разработка VLAN и IP-плана, определение ролей серверов, настройка эмуляции, проверка связности и резервного копирования."]),
        ("1 Общая часть", ["Объектом проектирования является офис компании «Ксанакс». Пользователи разделены на администрацию, бухгалтерию, отдел продаж и техническую службу.", "К сети предъявляются требования производительности, безопасности, управляемости, масштабируемости и резервного копирования. Для ограничения широковещательного трафика используется VLAN."]),
        ("2 Специальная часть", ["Физическая и логическая структуры проектируются раздельно. Физическая схема описывает соединения оборудования, логическая — VLAN, адреса, маршрутизацию и правила доступа.", "VLAN 10 используется администрацией, VLAN 20 бухгалтерией, VLAN 30 отделом продаж, VLAN 40 серверами, VLAN 50 управлением оборудованием и системой мониторинга.", "Серверные роли: SRV-01 — AD/DNS/DHCP; SRV-02 — файловый сервер; SRV-03 — резервное копирование; SRV-04 — FastAPI, SQLite и 1337 Control; SRV-05 — web/iOS gateway.", "Агент 1337 Control собирает CPU, RAM, диск, сеть, процессы, службы и HTTP-проверки. FastAPI принимает метрики с X-Agent-Token и сохраняет историю в SQLite."]),
        ("3 Практическая часть", ["В Cisco Packet Tracer размещаются маршрутизатор, L3-коммутатор, три access-коммутатора, 30 PC и 5 Server. Uplink-порты настраиваются trunk, пользовательские порты — access.", "На ядре создаются VLAN 10, 20, 30, 40 и 50, интерфейсы SVI с адресами шлюзов и включается ip routing. Рабочие станции получают адреса через DHCP, серверы используют статические адреса.", "Проверяются ping шлюзов, DHCP, DNS, доступ к файловому серверу, ограничения ACL и передача метрик на SRV-04. Скрипт backup_db.py создаёт копии SQLite с временной меткой и хранит последние 14 файлов."]),
        ("ЗАКЛЮЧЕНИЕ", ["Разработана ЛВС офиса «Ксанакс» на 30 рабочих мест и 5 серверов. Сформированы топология, VLAN и IP-план, определены серверные роли и порядок моделирования в Cisco Packet Tracer. Интеграция 1337 Control обеспечивает централизованный мониторинг технических узлов через web и iOS."]),
    ]
    for title, paragraphs in sections:
        doc.add_heading(title, level=1)
        for text in paragraphs:
            add_p(doc, text)
        if title == "2 Специальная часть" and (ASSETS / "topology.png").exists():
            doc.add_picture(str(ASSETS / "topology.png"), width=Cm(16))
            add_p(doc, "Рисунок 1 — Топология сети", center=True)

    doc.add_heading("СПИСОК СОКРАЩЕНИЙ", level=1)
    add_table(doc, ["Термин", "Значение"], [["ЛВС", "локальная вычислительная сеть"], ["VLAN", "виртуальная локальная сеть"], ["DHCP", "динамическая конфигурация узлов"], ["DNS", "служба доменных имён"], ["API", "программный интерфейс"], ["ACL", "список контроля доступа"]])
    doc.add_heading("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", level=1)
    for source in ["Кузин А.В. Компьютерные сети. — Москва: ИНФРА-М.", "Максимов Н.В. Компьютерные сети. — Москва: ИНФРА-М.", "Cisco Networking Academy. Материалы по коммутации и маршрутизации.", "FastAPI Documentation — https://fastapi.tiangolo.com/", "Python Documentation — https://docs.python.org/3/", "SQLite Documentation — https://sqlite.org/docs.html", "Apple SwiftUI Documentation — https://developer.apple.com/documentation/swiftui", "Исходный код проекта 1337 Control."]:
        add_p(doc, source)
    doc.save(OUT / "course_1337_control.docx")

def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides = [
        ("Разработка ЛВС офиса «Ксанакс»", "Курсовой проект\nСистема мониторинга 1337 Control\nКалачев Роман Дмитриевич\nУфа, 2026", None),
        ("Аннотация", "30 рабочих станций, 5 серверов, VLAN, Cisco Packet Tracer и мониторинг.", None),
        ("Актуальность", "Надёжность сети, безопасность и быстрое обнаружение отказов.", None),
        ("Цель и задачи", "Топология, адресация, VLAN, серверы, эмуляция и мониторинг.", None),
        ("Объект проекта", "Администрация — 6 ПК; бухгалтерия — 6; продажи — 12; техническая служба — 6.", None),
        ("Требования", "Производительность • безопасность • управляемость • резервирование.", None),
        ("Топология сети", "", ASSETS / "topology.png"),
        ("VLAN-сегментация", "VLAN 10 ADMIN\nVLAN 20 ACCOUNTING\nVLAN 30 SALES\nVLAN 40 SERVERS\nVLAN 50 MGMT", None),
        ("IP-план", "192.168.10.0/24 — ADMIN\n192.168.20.0/24 — ACCOUNTING\n192.168.30.0/24 — SALES\n192.168.40.0/24 — SERVERS\n192.168.50.0/24 — MGMT", None),
        ("Серверные роли", "", ASSETS / "server_roles.png"),
        ("1337 Control", "Агент → FastAPI API → SQLite → Web/iOS.", None),
        ("Собираемые метрики", "CPU, RAM, диск, сеть, процессы, службы, HTTP-проверки.", None),
        ("Cisco Packet Tracer", "R1, L3 core, 3 access-коммутатора, 30 PC и 5 Server.", None),
        ("Настройка VLAN", "SVI, ip routing, access-порты, trunk, DHCP и DNS.", None),
        ("Проверка сети", "Ping, DHCP, DNS, ACL, доступ к серверам и передача метрик.", None),
        ("Резервное копирование", "SQLite backup, временные метки, хранение последних 14 копий.", None),
        ("Темы интерфейса", "Cyber • Midnight • Light.", None),
        ("Результаты", "Спроектирована сеть и интегрирован мониторинг технических узлов.", None),
        ("Спасибо за внимание", "Вопросы?", None),
    ]
    for title, body, image in slides:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(10, 18, 32)
        box = slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(12.2), Inches(.7)).text_frame
        box.text = title; box.paragraphs[0].font.size = PPTPt(26); box.paragraphs[0].font.bold = True; box.paragraphs[0].font.color.rgb = RGBColor(60, 210, 235)
        if image and image.exists():
            slide.shapes.add_picture(str(image), Inches(.7), Inches(1.2), width=Inches(11.9), height=Inches(5.7))
        else:
            body_box = slide.shapes.add_textbox(Inches(.8), Inches(1.3), Inches(11.7), Inches(5.3)).text_frame
            body_box.text = body
            for paragraph in body_box.paragraphs:
                paragraph.font.size = PPTPt(20); paragraph.font.color.rgb = RGBColor(235, 242, 250)
    prs.save(OUT / "presentation_1337_control.pptx")

if __name__ == "__main__":
    build_docx(); build_pptx(); print("UTF8_REBUILT")
